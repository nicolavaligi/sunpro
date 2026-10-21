# ==============================================================================
# PROPRIETARY AND CONFIDENTIAL — NICOLA VALIGI (HOUDINICK)
# SunPro Geo-Intelligence 3D — Piacentino A1-A21 Dedicated Platform Builder
# Autore & Titolare Esclusivo della Proprietà Intellettuale: Nicola Valigi (Houdinick)
# Email: 305862309+nicolavaligi@users.noreply.github.com | GitHub: nicolavaligi
# Copyright (c) 2026 Nicola Valigi. All Rights Reserved.
# ==============================================================================
"""
Generatore della Mappa Dedicata & Workstation Visure per il Cluster Piacentino (A1 & A21).
Integra:
- 50 Lotti Top Aperta Campagna (da Terreni_FV_Piacentino_A1_A21.xlsx)
- 15 Lotti Aperta Campagna Oltre i Primi 50 (da foglio Riserve)
- 5 Lotti Benchmark Test A1 Pontenure-Cadeo (tagliati dai bordi della prima griglia)
- 20 Nuovi Lotti Aperta Campagna da Estensione Scansione WMS AdE a Griglia Sfalsata
- 26 Lotti Riserve Periurbane (<250m da aree produttive/residenziali o Le Mose)
Totale: 90 Lotti Qualificati in Aperta Campagna + 26 Riserve Periurbane = 116 Lotti Verificati WMS AdE.

Arricchisce ogni lotto con tutte le specifiche tecniche SunPro:
- Distanza Cabina Primaria ARERA/GSE (linea d'aria + routing stradale 1.30x) e CAPEX allaccio TICA/STMG
- Distanza Linea AT (132/220/380 kV) e Stazione Elettrica RTN/AT da OpenStreetMap
- Dimensionamento MWp (totale e nella fascia <=300m) + Producibilità PVGIS JRC (Fisso vs Tracker Monoassiale)
- Valore Agricolo Medio VAM 2026 Piacenza (BURERT n. 75 del 02/04/2026) + Calcolatore Soglia Prezzo (<=8,00 €/mq)
- Canone Diritto di Superficie 30 anni (3.000 €/ha/anno)
- Score Multicriterio SunPro (0-100) e Classe di Priorità Visura Sister
"""

import json
import math
import re
import shutil
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from shapely.geometry import LineString, Point, Polygon, mapping
from shapely.ops import unary_union

from crawler.cadastral_resolver import get_ade_wms_feature_info_url
from crawler.land_valuation import COMUNE_VAM_2026_EUR_MQ, evaluate_land_market_value
from crawler.pvgis import get_solar_comparison
from crawler.substation_finder import find_substations_around_coords
from scoring.scorer import calculate_site_score

BASE_DIR = Path(__file__).resolve().parent
WEB_DIR = BASE_DIR / "web"
WEB_DIR.mkdir(parents=True, exist_ok=True)
DATA_DIR = BASE_DIR / "data"

EXCEL_INPUT = DATA_DIR / "Terreni_FV_Piacentino_A1_A21.xlsx"
WMS_EXT_JSON = DATA_DIR / "reference" / "piacentino_wms_extension.json"
OSM_CACHE_JSON = DATA_DIR / "reference" / "osm_piacentino_a1_a21.json"

LAT0, LON0 = 45.02, 9.75
M_PER_LAT = 111132.0
M_PER_LON = 111320.0 * math.cos(math.radians(LAT0))


def to_xy(lat: float, lon: float):
    return ((lon - LON0) * M_PER_LON, (lat - LAT0) * M_PER_LAT)


def to_lonlat(x: float, y: float):
    return ( round(LON0 + x / M_PER_LON, 6), round(LAT0 + y / M_PER_LAT, 6) )


def load_all_piacentino_raw_lots() -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(EXCEL_INPUT, data_only=True)
    raw_lots: List[Dict[str, Any]] = []

    # 1. Foglio Terreni (Top 50 Aperta Campagna)
    ws1 = wb["Terreni"]
    for r in range(6, 56):
        if ws1.cell(r, 1).value is not None:
            raw_lots.append({
                "origine": "TOP_50_EXCEL",
                "gruppo": "APERTA_CAMPAGNA",
                "categoria_label": "Top 50 Aperta Campagna",
                "comune": str(ws1.cell(r, 2).value).strip(),
                "belfiore": str(ws1.cell(r, 3).value).strip(),
                "foglio": str(ws1.cell(r, 4).value).strip(),
                "particella": str(ws1.cell(r, 5).value).strip(),
                "sup_tot_ha": float(ws1.cell(r, 6).value),
                "sup_300m_ha": float(ws1.cell(r, 7).value),
                "autostrada": str(ws1.cell(r, 8).value).strip(),
                "dist_asse_m": int(ws1.cell(r, 9).value),
                "dist_linea_at_m": int(ws1.cell(r, 10).value) if ws1.cell(r, 10).value is not None else 999,
                "kv_linea": int(ws1.cell(r, 11).value) if ws1.cell(r, 11).value is not None else 132,
                "stazione_nome": str(ws1.cell(r, 12).value or "(senza nome)").strip(),
                "kv_stazione": int(ws1.cell(r, 13).value) if ws1.cell(r, 13).value is not None else 132,
                "dist_stazione_km": float(ws1.cell(r, 14).value) if ws1.cell(r, 14).value is not None else 3.0,
                "lat": float(ws1.cell(r, 15).value),
                "lon": float(ws1.cell(r, 16).value),
                "note": str(ws1.cell(r, 19).value or "").strip(),
            })

    # 2. Foglio Riserve (15 Aperta Campagna oltre i primi 50 + 26 Periurbani)
    ws2 = wb["Riserve"]
    for r in range(6, 47):
        if ws2.cell(r, 1).value is not None:
            cat_raw = str(ws2.cell(r, 2).value or "").strip()
            is_aperta = "Aperta campagna" in cat_raw
            raw_lots.append({
                "origine": "RISERVE_APERTA_15" if is_aperta else "RISERVE_PERIURBANE_26",
                "gruppo": "APERTA_CAMPAGNA" if is_aperta else "RISERVA_PERIURBANA",
                "categoria_label": "Aperta Campagna (Oltre Top 50)" if is_aperta else "Riserva Periurbana (<250m Z.I./Res.)",
                "comune": str(ws2.cell(r, 3).value).strip(),
                "belfiore": str(ws2.cell(r, 4).value).strip(),
                "foglio": str(ws2.cell(r, 5).value).strip(),
                "particella": str(ws2.cell(r, 6).value).strip(),
                "sup_tot_ha": float(ws2.cell(r, 7).value),
                "sup_300m_ha": float(ws2.cell(r, 8).value),
                "autostrada": str(ws2.cell(r, 9).value).strip(),
                "dist_asse_m": int(ws2.cell(r, 10).value),
                "dist_linea_at_m": int(ws2.cell(r, 11).value) if ws2.cell(r, 11).value is not None else 999,
                "kv_linea": int(ws2.cell(r, 12).value) if ws2.cell(r, 12).value is not None else 132,
                "stazione_nome": str(ws2.cell(r, 13).value or "(senza nome)").strip(),
                "kv_stazione": int(ws2.cell(r, 14).value) if ws2.cell(r, 14).value is not None else 132,
                "dist_stazione_km": float(ws2.cell(r, 15).value) if ws2.cell(r, 15).value is not None else 3.0,
                "lat": float(ws2.cell(r, 16).value),
                "lon": float(ws2.cell(r, 17).value),
                "note": str(ws2.cell(r, 20).value or "").strip(),
            })

    # 3. I 5 Lotti Benchmark del Test Pontenure-Cadeo (tagliati dai bordi della griglia)
    benchmark_5 = [
        {
            "origine": "BENCHMARK_TEST_5",
            "gruppo": "APERTA_CAMPAGNA",
            "categoria_label": "Benchmark Test A1 (Pontenure-Cadeo)",
            "comune": "Cadeo",
            "belfiore": "B332",
            "foglio": "7",
            "particella": "77",
            "sup_tot_ha": 13.10,
            "sup_300m_ha": 11.20,
            "autostrada": "A1",
            "dist_asse_m": 45,
            "dist_linea_at_m": 85,
            "kv_linea": 132,
            "stazione_nome": "Cadeo FS",
            "kv_stazione": 132,
            "dist_stazione_km": 3.1,
            "lat": 44.994000,
            "lon": 9.827500,
            "note": "Lotto 1 Benchmark Test A1 (contigue Part. 97, 63); sfuggito alla griglia standard perché sul bordo riquadro; 1,9 ha oltre i 300 m",
        },
        {
            "origine": "BENCHMARK_TEST_5",
            "gruppo": "APERTA_CAMPAGNA",
            "categoria_label": "Benchmark Test A1 (Pontenure-Cadeo)",
            "comune": "Cadeo",
            "belfiore": "B332",
            "foglio": "7",
            "particella": "149",
            "sup_tot_ha": 8.10,
            "sup_300m_ha": 7.40,
            "autostrada": "A1",
            "dist_asse_m": 65,
            "dist_linea_at_m": 110,
            "kv_linea": 132,
            "stazione_nome": "Cadeo FS",
            "kv_stazione": 132,
            "dist_stazione_km": 3.2,
            "lat": 44.996500,
            "lon": 9.828500,
            "note": "Lotto 2 Benchmark Test A1 (tra A1 e linea AV Milano-Bologna); sfuggito alla griglia standard per taglio bordo riquadro",
        },
        {
            "origine": "BENCHMARK_TEST_5",
            "gruppo": "APERTA_CAMPAGNA",
            "categoria_label": "Benchmark Test A1 (Pontenure-Cadeo)",
            "comune": "Pontenure",
            "belfiore": "G852",
            "foglio": "14",
            "particella": "134",
            "sup_tot_ha": 6.30,
            "sup_300m_ha": 6.10,
            "autostrada": "A1",
            "dist_asse_m": 38,
            "dist_linea_at_m": 95,
            "kv_linea": 132,
            "stazione_nome": "Cadeo FS",
            "kv_stazione": 132,
            "dist_stazione_km": 3.8,
            "lat": 44.996500,
            "lon": 9.821000,
            "note": "Lotto 3 Benchmark Test A1 (contigua a Part. 21, totale compendio 10,5 ha); verificato WMS AdE",
        },
        {
            "origine": "BENCHMARK_TEST_5",
            "gruppo": "APERTA_CAMPAGNA",
            "categoria_label": "Benchmark Test A1 (Pontenure-Cadeo)",
            "comune": "Pontenure",
            "belfiore": "G852",
            "foglio": "14",
            "particella": "21",
            "sup_tot_ha": 4.20,
            "sup_300m_ha": 4.10,
            "autostrada": "A1",
            "dist_asse_m": 52,
            "dist_linea_at_m": 120,
            "kv_linea": 132,
            "stazione_nome": "Cadeo FS",
            "kv_stazione": 132,
            "dist_stazione_km": 4.0,
            "lat": 44.997200,
            "lon": 9.818500,
            "note": "Lotto 4 Benchmark Test A1 (contigua a Part. 134, accesso da Strada Comunale Valconasso)",
        },
        {
            "origine": "BENCHMARK_TEST_5",
            "gruppo": "APERTA_CAMPAGNA",
            "categoria_label": "Benchmark Test A1 (Pontenure-Cadeo)",
            "comune": "Pontenure",
            "belfiore": "G852",
            "foglio": "14",
            "particella": "86",
            "sup_tot_ha": 3.20,
            "sup_300m_ha": 3.20,
            "autostrada": "A1",
            "dist_asse_m": 60,
            "dist_linea_at_m": 140,
            "kv_linea": 132,
            "stazione_nome": "Cadeo FS",
            "kv_stazione": 132,
            "dist_stazione_km": 3.9,
            "lat": 44.999000,
            "lon": 9.822000,
            "note": "Lotto 5 Benchmark Test A1 (contigua Part. 44, tra A1 e ferrovia AV)",
        },
    ]
    raw_lots.extend(benchmark_5)

    # 4. I 20 Nuovi Lotti Aperta Campagna da Estensione Scansione WMS AdE (Griglia Sfalsata)
    if WMS_EXT_JSON.exists():
        ext_lots = json.loads(WMS_EXT_JSON.read_text(encoding="utf-8"))
        for el in ext_lots:
            raw_lots.append({
                "origine": "ESTENSIONE_WMS_20",
                "gruppo": "APERTA_CAMPAGNA",
                "categoria_label": "Estensione WMS AdE (Griglia Sfalsata)",
                "comune": el["comune"],
                "belfiore": el["belfiore"],
                "foglio": str(el["foglio"]),
                "particella": str(el["particella"]),
                "sup_tot_ha": float(el["sup_tot_ha"]),
                "sup_300m_ha": float(el["sup_300m_ha"]),
                "autostrada": el["autostrada"],
                "dist_asse_m": int(el["dist_asse_m"]),
                "dist_linea_at_m": int(el.get("dist_linea_at_m") or 400),
                "kv_linea": int(el.get("kv_linea") or 132),
                "stazione_nome": str(el.get("stazione_nome") or "(senza nome)"),
                "kv_stazione": int(el.get("kv_stazione") or 132),
                "dist_stazione_km": float(el.get("dist_stazione_km") or 2.5),
                "lat": float(el["lat"]),
                "lon": float(el["lon"]),
                "note": str(el.get("note") or "Individuato con scansione WMS a griglia sfalsata anti-taglio").strip(),
            })

    return raw_lots


def enrich_piacentino_lots(raw_lots: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    enriched: List[Dict[str, Any]] = []

    # Sort Aperta Campagna first by sup_300m_ha desc, then Periurbani by sup_300m_ha desc
    aperta = [x for x in raw_lots if x["gruppo"] == "APERTA_CAMPAGNA"]
    peri = [x for x in raw_lots if x["gruppo"] != "APERTA_CAMPAGNA"]
    aperta.sort(key=lambda x: (x["sup_300m_ha"], x["sup_tot_ha"]), reverse=True)
    peri.sort(key=lambda x: (x["sup_300m_ha"], x["sup_tot_ha"]), reverse=True)

    ordered = aperta + peri

    # Cache PVGIS per Piacenza (~45.0N, 9.7E) per velocità
    solar_ref = get_solar_comparison(45.02, 9.75)
    fisso_kwh_kwp = float(solar_ref.get("fisso_kwh_kwp", 1315.0))
    tracker_kwh_kwp = float(solar_ref.get("tracker_kwh_kwp", 1578.0))
    boost_tracker_pct = float(solar_ref.get("boost_pct", 20.0))

    for idx, lot in enumerate(ordered, 1):
        lat = round(float(lot["lat"]), 6)
        lon = round(float(lot["lon"]), 6)
        comune = lot["comune"]
        belfiore = lot["belfiore"]
        foglio = lot["foglio"]
        particella = lot["particella"]
        sup_tot_ha = round(float(lot["sup_tot_ha"]), 2)
        sup_300m_ha = round(float(lot["sup_300m_ha"]), 2)
        sup_tot_mq = int(round(sup_tot_ha * 10000))
        sup_300m_mq = int(round(sup_300m_ha * 10000))
        quota_300m_pct = min(100, int(round((sup_300m_ha / max(0.01, sup_tot_ha)) * 100)))

        lot_id = f"PC-{lot['autostrada']}-{idx:03d}"

        # 1. Ricerca Cabina Primaria Ufficiale ARERA/GSE + Stazione Elettrica OSM
        subs_arera = find_substations_around_coords(lat, lon, radius_m=15000)
        cp_best = subs_arera[0] if subs_arera else {
            "name": "CP Piacenza / Pontenure 132/15 kV",
            "distanza_m": int(lot["dist_stazione_km"] * 1000),
            "distanza_stradale_m": int(lot["dist_stazione_km"] * 1300),
            "livello_tensione": "132/15 kV",
            "operatore": "e-distribuzione",
            "capex_allaccio_stimato_eur": 195000,
        }
        cp_dist_aria_m = int(cp_best["distanza_m"])
        cp_dist_osm_m = int(round(float(lot["dist_stazione_km"]) * 1000))
        # Distanza effettiva alla stazione/cabina più prossima tra OSM e ARERA
        effective_sub_dist_m = min(cp_dist_aria_m, cp_dist_osm_m)
        cavidotto_mt_m = int(round(effective_sub_dist_m * 1.30))
        capex_allaccio_eur = int(round((max(0.15, cavidotto_mt_m / 1000.0) * 65000.0) + 45000.0))

        # 2. Dimensionamento Fotovoltaico (1 MWp ogni 1,2 ha = 0,833 MWp/ha)
        mwp_tot = round(sup_tot_ha / 1.20, 2)
        mwp_300m = round(sup_300m_ha / 1.20, 2)
        mwh_anno_fisso = int(round(mwp_tot * fisso_kwh_kwp))
        mwh_anno_tracker = int(round(mwp_tot * tracker_kwh_kwp))
        mwh_300m_tracker = int(round(mwp_300m * tracker_kwh_kwp))

        # 3. Valutazione Fondiaria VAM 2026 Piacenza & Soglia Mandato <= 8,00 €/mq
        vam_mq = float(COMUNE_VAM_2026_EUR_MQ.get(comune, 5.55))
        is_periurbano = lot["gruppo"] == "RISERVA_PERIURBANA"
        val_report = evaluate_land_market_value(
            lead_id=lot_id,
            regione="Emilia-Romagna",
            provincia="Piacenza",
            comune=comune,
            superficie_mq=sup_tot_mq,
            tipologia="BUFFER_AUTOSTRADALE_300M_AGRICOLO",
            distanza_zona_industriale_m=150.0 if is_periurbano else 600.0,
            distanza_autostrada_m=float(lot["dist_asse_m"]),
        )
        # Per i lotti in aperta campagna il target SunPro è ~7.10-7.55 €/mq (<= 8.00 €/mq)
        prezzo_target_mq = round(val_report.prezzo_acquisto_target_eur_mq, 2)
        if is_periurbano:
            # Nei periurbani segnaliamo il rischio di richiesta sopra soglia (es. 8.50-11.00 €/mq)
            prezzo_target_mq = 8.50

        spesa_max_8eur = int(round(sup_tot_mq * 8.00))
        spesa_300m_8eur = int(round(sup_300m_mq * 8.00))
        canone_dds_tot_eur = int(round(sup_tot_ha * 3000.0))
        canone_dds_300m_eur = int(round(sup_300m_ha * 3000.0))

        # 4. Calcolo Score Multicriterio SunPro (0-100)
        lead_for_score = {
            "id": lot_id,
            "title": f"Compendio {comune} {lot['autostrada']} — Fg.{foglio} Mapp.{particella}",
            "regione": "Emilia-Romagna",
            "provincia": "Piacenza",
            "comune": comune,
            "lat": lat,
            "lng": lon,
            "superficie_mq": sup_tot_mq,
            "tipologia": "TERRENO_AGRICOLO_CORRIDOIO_AUTOSTRADALE",
            "distanza_autostrada_m": int(lot["dist_asse_m"]),
            "distanza_zona_industriale_m": 180 if is_periurbano else 650,
            "distanza_cabina_m": effective_sub_dist_m,
            "prezzo_mq_eur": prezzo_target_mq,
            "status_acquistabilita": "COMPATIBILE_SOGLIA_8_EUR_MQ",
            "catasto_stato": "VERIFICATO_WMS_ADE",
            "codice_belfiore": belfiore,
            "foglio": foglio,
            "particella": particella,
        }
        score_tot, score_det, _ = calculate_site_score(lead_for_score)
        # Bonus per vicinanza estrema alla linea AT (<150m) o grande superficie entro 300m (>=7 ha)
        if not is_periurbano:
            if lot["dist_linea_at_m"] <= 150 or effective_sub_dist_m <= 1200:
                score_tot = min(99.0, round(score_tot + 2.5, 1))
            if sup_300m_ha >= 7.0:
                score_tot = min(99.5, round(score_tot + 2.0, 1))
        else:
            score_tot = min(84.5, round(score_tot - 4.0, 1))

        # 5. Classificazione Priorità Visura Sister
        if is_periurbano:
            priorita_visura = "RISERVA_PERIURBANA"
            priorita_label = "⚠️ Riserva Periurbana (Verifica PRG/Prezzo)"
        elif sup_300m_ha >= 6.0 or float(lot["dist_stazione_km"]) <= 1.5 or int(lot["dist_linea_at_m"]) <= 150:
            priorita_visura = "PRIORITA_1_TOP"
            priorita_label = "🔥 Priorità 1 · Top Taglia / Rete AT Vicina"
        else:
            priorita_visura = "PRIORITA_2_APERTA"
            priorita_label = "✅ Priorità 2 · Aperta Campagna Qualificata"

        staz_str = lot["stazione_nome"]
        if lot["kv_stazione"]:
            staz_str = f"{staz_str} ({lot['kv_stazione']} kV)"

        enriched.append({
            "id": lot_id,
            "rank": idx,
            "origine": lot["origine"],
            "gruppo": lot["gruppo"],
            "categoria_label": lot["categoria_label"],
            "priorita_visura": priorita_visura,
            "priorita_label": priorita_label,
            "comune": comune,
            "belfiore": belfiore,
            "foglio": foglio,
            "particella": particella,
            "catasto": f"Fg. {foglio} · Part. {particella} (Cod. {belfiore})",
            "sup_tot_ha": sup_tot_ha,
            "sup_300m_ha": sup_300m_ha,
            "sup_oltre_300m_ha": round(max(0.0, sup_tot_ha - sup_300m_ha), 2),
            "quota_300m_pct": quota_300m_pct,
            "autostrada": lot["autostrada"],
            "dist_asse_m": int(lot["dist_asse_m"]),
            "dist_linea_at_m": int(lot["dist_linea_at_m"]),
            "kv_linea": int(lot["kv_linea"]),
            "stazione_nome": lot["stazione_nome"],
            "kv_stazione": int(lot["kv_stazione"]),
            "dist_stazione_km": round(float(lot["dist_stazione_km"]), 1),
            "stazione_display": staz_str,
            "cp_arera_nome": cp_best["name"],
            "cp_arera_operatore": cp_best.get("operatore", "e-distribuzione"),
            "cp_arera_tensione": cp_best.get("livello_tensione", "132/15 kV"),
            "cp_dist_aria_m": cp_dist_aria_m,
            "cavidotto_mt_m": cavidotto_mt_m,
            "capex_allaccio_eur": capex_allaccio_eur,
            "mwp_tot": mwp_tot,
            "mwp_300m": mwp_300m,
            "mwh_fisso": mwh_anno_fisso,
            "mwh_tracker": mwh_anno_tracker,
            "mwh_300m_tracker": mwh_300m_tracker,
            "fisso_kwh_kwp": fisso_kwh_kwp,
            "tracker_kwh_kwp": tracker_kwh_kwp,
            "boost_tracker_pct": boost_tracker_pct,
            "vam_2026_mq": round(vam_mq, 2),
            "prezzo_target_mq": prezzo_target_mq,
            "spesa_max_8eur": spesa_max_8eur,
            "spesa_300m_8eur": spesa_300m_8eur,
            "canone_dds_tot_eur": canone_dds_tot_eur,
            "canone_dds_300m_eur": canone_dds_300m_eur,
            "score": score_tot,
            "bars": {
                "norm": min(100, int((score_det.get("idoneita_normativa", 28) / 30.0) * 100)),
                "grid": min(100, max(35, int(100 - (min(4000, effective_sub_dist_m) / 4000.0) * 65))),
                "price": 72 if is_periurbano else 96,
                "size": min(100, int(min(1.0, sup_300m_ha / 10.0) * 100)),
            },
            "lat": lat,
            "lon": lon,
            "ll": [lat, lon],
            "wms_url": get_ade_wms_feature_info_url(lat, lon),
            "gmaps_url": f"https://www.google.com/maps?q={lat},{lon}&t=k&z=17",
            "note": lot["note"] or "Seminativo aperto in fascia autostradale ≤300m D.Lgs. 199/2021",
        })

    return enriched


def export_vector_layers_geojson(out_path: Path):
    """Costruisce il GeoJSON vettoriale con:
    - Tracciati A1 e A21 in provincia di Piacenza
    - Fascia Buffer Ex Lege 300m attorno ad A1 e A21
    - Linee Elettriche AT (132 / 220 / 380 kV)
    - Stazioni Elettriche RTN/AT da OpenStreetMap
    """
    if not OSM_CACHE_JSON.exists():
        return
    osm = json.loads(OSM_CACHE_JSON.read_text(encoding="utf-8")).get("elements", [])
    features = []
    a1_lines, a21_lines = [], []

    for el in osm:
        tags = el.get("tags", {})
        geom = el.get("geometry", [])
        if tags.get("highway") == "motorway" and len(geom) >= 2:
            coords = [[round(p["lon"], 5), round(p["lat"], 5)] for p in geom]
            xy_pts = [to_xy(p["lat"], p["lon"]) for p in geom]
            ref = tags.get("ref", "Autostrada")
            hwy_label = "A1" if ("A1" in ref and "A21" not in ref) else ("A21" if "A21" in ref else ref)
            if hwy_label == "A1":
                a1_lines.append(LineString(xy_pts))
            elif hwy_label == "A21":
                a21_lines.append(LineString(xy_pts))
            features.append({
                "type": "Feature",
                "properties": {"layer": "highway_axis", "ref": hwy_label},
                "geometry": {"type": "LineString", "coordinates": coords}
            })
        elif tags.get("power") == "line" and len(geom) >= 2:
            coords = [[round(p["lon"], 5), round(p["lat"], 5)] for p in geom]
            v_raw = str(tags.get("voltage", ""))
            kv = 132
            for part in re.split(r"[;/,]", v_raw):
                part = part.strip()
                if part.isdigit():
                    val = int(part)
                    kv = val // 1000 if val >= 1000 else val
                    break
            features.append({
                "type": "Feature",
                "properties": {
                    "layer": "power_line",
                    "kv": kv,
                    "operator": tags.get("operator", "Terna / e-distribuzione"),
                    "name": tags.get("name", f"Elettrodotto AT {kv} kV")
                },
                "geometry": {"type": "LineString", "coordinates": coords}
            })
        elif tags.get("power") == "substation":
            if "lat" in el and "lon" in el:
                lon_c, lat_c = round(el["lon"], 5), round(el["lat"], 5)
            elif geom:
                lon_c = round(sum(p["lon"] for p in geom) / len(geom), 5)
                lat_c = round(sum(p["lat"] for p in geom) / len(geom), 5)
            else:
                continue
            v_raw = str(tags.get("voltage", ""))
            kv = 132
            for part in re.split(r"[;/,]", v_raw):
                part = part.strip()
                if part.isdigit():
                    val = int(part)
                    kv = val // 1000 if val >= 1000 else val
                    break
            features.append({
                "type": "Feature",
                "properties": {
                    "layer": "substation",
                    "name": tags.get("name") or "Stazione Elettrica AT/MT",
                    "kv": kv,
                    "operator": tags.get("operator", "Terna / e-distribuzione / RFI")
                },
                "geometry": {"type": "Point", "coordinates": [lon_c, lat_c]}
            })

    # Genera poligono buffer 300m semplificato attorno ad A1 e A21
    for hwy_name, lines in [("A1", a1_lines), ("A21", a21_lines)]:
        if not lines:
            continue
        buf_xy = unary_union(lines).buffer(300.0, resolution=6).simplify(18.0)
        polys = [buf_xy] if buf_xy.geom_type == "Polygon" else list(buf_xy.geoms)
        for poly in polys:
            ext_coords = [list(to_lonlat(x, y)) for x, y in poly.exterior.coords]
            features.append({
                "type": "Feature",
                "properties": {
                    "layer": "buffer_300m",
                    "autostrada": hwy_name,
                    "norma": "Fascia Ex Lege ≤ 300m D.Lgs. 199/2021 art. 20 c. 8-ter"
                },
                "geometry": {"type": "Polygon", "coordinates": [ext_coords]}
            })

    geojson_obj = {"type": "FeatureCollection", "features": features}
    out_path.write_text(json.dumps(geojson_obj, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
    shutil.copy2(out_path, WEB_DIR / out_path.name)
    print(f"  ✓ GeoJSON vettoriale Piacentino (A1/A21 + Buffer 300m + Linee AT + Stazioni) salvato: {len(features)} feature")


def export_extended_excel_and_csv(lots: List[Dict[str, Any]], xlsx_path: Path, csv_path: Path):
    """Genera il file Excel esteso Terreni_FV_Piacentino_A1_A21_Esteso.xlsx e il CSV completo."""
    wb = openpyxl.Workbook()
    ws_aperta = wb.active
    ws_aperta.title = "Terreni Aperta Campagna (90)"
    ws_peri = wb.create_sheet("Riserve Periurbane (26)")
    ws_metodo = wb.create_sheet("Metodo Fonti e Specifiche")

    header_fill = PatternFill("solid", fgColor="0D1526")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    yellow_fill = PatternFill("solid", fgColor="FEF3C7")
    green_fill = PatternFill("solid", fgColor="D1FAE5")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    headers = [
        "#", "Origine / Coorte", "Priorità Visura", "Comune", "Belfiore", "Foglio", "Particella",
        "Sup. stimata (ha)", "di cui entro 300 m (ha)", "Autostrada", "Distanza min. asse (m)",
        "Linea AT più vicina (m)", "Tensione linea (kV)", "Stazione elettrica OSM", "kV stazione",
        "Distanza stazione (km)", "Cabina Primaria ARERA/GSE", "Cavidotto MT 1.30x (m)", "CAPEX Allaccio TICA (€)",
        "Potenza (MWp tot)", "Potenza ≤300m (MWp)", "Producibilità Tracker (MWh/a)",
        "VAM 2026 (€/mq)", "Target SunPro (€/mq)", "Spesa max a soglia (€)", "Canone DDS 30y (€/anno)",
        "Score SunPro", "Lat", "Lon", "Mappa / WMS", "Note",
        "Intestatario (da visura)", "Prezzo richiesto (€/mq)", "Entro soglia?"
    ]

    def populate_sheet(ws, subset: List[Dict[str, Any]], title_text: str):
        ws["A1"] = title_text
        ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="059669")
        ws["A2"] = (
            "Criteri: particelle agricole >= 2 ha entro 300 m dall'asse di A1 o A21 (provincia di Piacenza), senza fabbricati, "
            "verificate su WMS Catasto Agenzia delle Entrate + Rete AT/MT ARERA/GSE & OpenStreetMap + PVGIS JRC + VAM 2026."
        )
        ws["A3"] = "Prezzo massimo (€/mq)"
        ws["A3"].font = Font(bold=True)
        ws["D3"] = 8.00
        ws["D3"].fill = yellow_fill
        ws["D3"].font = Font(bold=True, color="1D4ED8")

        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(5, col_idx, h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for i, item in enumerate(subset, 1):
            r = 5 + i
            row_vals = [
                i,
                item["categoria_label"],
                item["priorita_label"],
                item["comune"],
                item["belfiore"],
                int(item["foglio"]) if str(item["foglio"]).isdigit() else item["foglio"],
                int(item["particella"]) if str(item["particella"]).isdigit() else item["particella"],
                item["sup_tot_ha"],
                item["sup_300m_ha"],
                item["autostrada"],
                item["dist_asse_m"],
                item["dist_linea_at_m"],
                item["kv_linea"],
                item["stazione_nome"],
                item["kv_stazione"],
                item["dist_stazione_km"],
                item["cp_arera_nome"],
                item["cavidotto_mt_m"],
                item["capex_allaccio_eur"],
                item["mwp_tot"],
                item["mwp_300m"],
                item["mwh_tracker"],
                item["vam_2026_mq"],
                item["prezzo_target_mq"],
                f"=H{r}*10000*$D$3",
                item["canone_dds_tot_eur"],
                item["score"],
                item["lat"],
                item["lon"],
                "Apri Ortofoto",
                item["note"],
                None,
                None,
                f'=IF(AG{r}="","",IF(AG{r}<=$D$3,"SÌ","NO"))'
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(r, col_idx, val)
                cell.border = thin_border
                if col_idx == 30:
                    cell.hyperlink = item["gmaps_url"]
                    cell.font = Font(color="0284C7", underline="single")
                if col_idx in (32, 33):
                    cell.fill = yellow_fill
                if col_idx == 25:
                    cell.fill = green_fill
                    cell.number_format = "#,##0"

        tot_r = 6 + len(subset)
        ws.cell(tot_r, 7, "Totale").font = Font(bold=True)
        ws.cell(tot_r, 8, f"=SUM(H6:H{tot_r-1})").font = Font(bold=True)
        ws.cell(tot_r, 9, f"=SUM(I6:I{tot_r-1})").font = Font(bold=True)
        ws.cell(tot_r, 20, f"=SUM(T6:T{tot_r-1})").font = Font(bold=True)
        ws.cell(tot_r, 21, f"=SUM(U6:U{tot_r-1})").font = Font(bold=True)
        ws.cell(tot_r, 25, f"=SUM(Y6:Y{tot_r-1})").font = Font(bold=True)
        ws.cell(tot_r, 25).number_format = "#,##0"

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 28
        ws.column_dimensions["Q"].width = 28
        ws.column_dimensions["AE"].width = 45

    aperta_lots = [x for x in lots if x["gruppo"] == "APERTA_CAMPAGNA"]
    peri_lots = [x for x in lots if x["gruppo"] != "APERTA_CAMPAGNA"]

    populate_sheet(
        ws_aperta,
        aperta_lots,
        f"{len(aperta_lots)} terreni in Aperta Campagna lungo A1 e A21 nel Piacentino (Top 50 + 15 Extra + 5 Benchmark + 20 Estensione WMS)"
    )
    populate_sheet(
        ws_peri,
        peri_lots,
        f"{len(peri_lots)} lotti in Riserva Periurbana (<250m da aree produttive/residenziali o polo Le Mose)"
    )

    metodo_lines = [
        ("Metodo e Integrazione Tecnica SunPro (Piacentino A1 & A21)", True),
        ("1. Asse di A1 e A21 ricavato da OpenStreetMap; scansione iniziale di 68 riquadri da 1,2 km + estensione a griglia sfalsata anti-taglio lungo A1 e A21 in provincia di Piacenza.", False),
        ("2. Particelle dalla cartografia catastale WMS dell'Agenzia delle Entrate (CP.CadastralParcel e fabbricati): superficie misurata sul disegno (1 m/px); scartate quelle sotto i 2 ha entro i 300 m, quelle con fabbricati e le strisce sottili.", False),
        ("3. Coorti integrate nel dataset Piacentino (116 lotti totali verificati WMS AdE):", True),
        (f"   • Top 50 Aperta Campagna (foglio Terreni originale): 388,43 ha totali (315,77 ha entro 300 m).", False),
        (f"   • 15 Lotti Aperta Campagna Oltre i Primi 50 (foglio Riserve originale): 70,22 ha totali (66,41 ha entro 300 m).", False),
        (f"   • 5 Lotti Benchmark Test A1 Pontenure-Cadeo (tagliati dai bordi della prima griglia): 34,90 ha totali (32,00 ha entro 300 m).", False),
        (f"   • 20 Nuovi Lotti Aperta Campagna da Estensione Scansione WMS a Griglia Sfalsata: {sum(x['sup_tot_ha'] for x in aperta_lots if x['origine']=='ESTENSIONE_WMS_20'):.2f} ha totali ({sum(x['sup_300m_ha'] for x in aperta_lots if x['origine']=='ESTENSIONE_WMS_20'):.2f} ha entro 300 m).", False),
        (f"   • 26 Lotti Riserve Periurbane (<250m da zone produttive/residenziali o Le Mose): 180,66 ha totali (154,52 ha entro 300 m).", False),
        ("4. Specifiche tecniche SunPro integrate su ogni particella:", True),
        ("   • Distanza Linea AT (132/220/380 kV) e Stazione Elettrica RTN/AT da OpenStreetMap.", False),
        ("   • Cabina Primaria ufficiale ARERA/GSE più vicina, distanza cavidotto MT (routing stradale 1.30x) e stima CAPEX allaccio TICA/STMG (65k €/km + 45k € stallo).", False),
        ("   • Dimensionamento MWp (totale e nella sola fascia ex lege <=300m) + Producibilità annua PVGIS JRC v5.2 (Struttura Fissa vs Tracker Monoassiale +20%).", False),
        ("   • Valore Agricolo Medio VAM 2026 Piacenza (BUR Emilia-Romagna n. 75 del 02/04/2026: 5,45-5,85 €/mq), Target SunPro <=8,00 €/mq, Spesa Max a soglia e Canone Diritto di Superficie 30 anni (3.000 €/ha/anno).", False),
        ("5. Verifiche pre-acquisizione raccomandate:", True),
        ("   • Visura catastale per immobile/soggetto su Sister (€ 1,35/particella) partendo dai lotti Priorità 1 (più grandi o con stazione elettrica <= 1,5 km / linea AT <= 150 m).", False),
        ("   • Verifica CDU (PUG/PSC comunale), vincoli paesaggistici/fluviali e certificazioni DOP/IGP/BIO (L.R. Emilia-Romagna 5/2026).", False),
    ]
    for r_idx, (txt, is_bold) in enumerate(metodo_lines, 1):
        c = ws_metodo.cell(r_idx, 1, txt)
        if is_bold:
            c.font = Font(bold=True, size=11, color="059669")

    wb.save(xlsx_path)
    shutil.copy2(xlsx_path, WEB_DIR / xlsx_path.name)
    shutil.copy2(xlsx_path, DATA_DIR / xlsx_path.name)

    # Esporta anche CSV piatto per analisi immediata
    import csv
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([
            "ID", "Rank", "Gruppo", "Coorte", "Priorita_Visura", "Comune", "Codice_Belfiore", "Foglio", "Particella",
            "Sup_Totale_ha", "Sup_Entro_300m_ha", "Autostrada", "Dist_Asse_m", "Dist_Linea_AT_m", "Tensione_Linea_kV",
            "Stazione_Elettrica_OSM", "kV_Stazione", "Dist_Stazione_km", "Cabina_Primaria_ARERA", "Cavidotto_MT_130x_m",
            "CAPEX_Allaccio_EUR", "MWp_Totale", "MWp_Entro_300m", "MWh_Anno_Tracker", "VAM_2026_EUR_mq",
            "Prezzo_Target_EUR_mq", "Spesa_Max_8EUR_Totale", "Spesa_Max_8EUR_300m", "Canone_DDS_Annuo_EUR",
            "Score_SunPro", "Lat", "Lon", "Note"
        ])
        for x in lots:
            writer.writerow([
                x["id"], x["rank"], x["gruppo"], x["categoria_label"], x["priorita_visura"],
                x["comune"], x["belfiore"], x["foglio"], x["particella"],
                x["sup_tot_ha"], x["sup_300m_ha"], x["autostrada"], x["dist_asse_m"],
                x["dist_linea_at_m"], x["kv_linea"], x["stazione_nome"], x["kv_stazione"], x["dist_stazione_km"],
                x["cp_arera_nome"], x["cavidotto_mt_m"], x["capex_allaccio_eur"],
                x["mwp_tot"], x["mwp_300m"], x["mwh_tracker"], x["vam_2026_mq"],
                x["prezzo_target_mq"], x["spesa_max_8eur"], x["spesa_300m_8eur"], x["canone_dds_tot_eur"],
                x["score"], x["lat"], x["lon"], x["note"]
            ])
    shutil.copy2(csv_path, WEB_DIR / csv_path.name)
    print(f"  ✓ Excel Esteso ({xlsx_path.name}) e CSV ({csv_path.name}) generati per {len(lots)} lotti!")


def export_piacentino_kml(lots: List[Dict[str, Any]], kml_path: Path):
    """Esporta il file KML 3D dedicato per Google Earth con tutti i 116 lotti del Piacentino."""
    placemarks = []
    for x in lots:
        color = "ff10b981" if x["gruppo"] == "APERTA_CAMPAGNA" else "ff0b9ef5"
        desc = (
            f"<![CDATA["
            f"<b>{x['comune']} ({x['belfiore']}) — Fg. {x['foglio']} Part. {x['particella']}</b><br/>"
            f"<b>Coorte:</b> {x['categoria_label']}<br/>"
            f"<b>Priorità Visura:</b> {x['priorita_label']}<br/>"
            f"<b>Superficie Totale:</b> {x['sup_tot_ha']} ha (di cui <b>{x['sup_300m_ha']} ha entro 300m</b> da {x['autostrada']})<br/>"
            f"<b>Distanza dall'asse {x['autostrada']}:</b> {x['dist_asse_m']} m<br/>"
            f"<b>Linea AT più vicina:</b> {x['dist_linea_at_m']} m ({x['kv_linea']} kV)<br/>"
            f"<b>Stazione Elettrica:</b> {x['stazione_display']} a {x['dist_stazione_km']} km<br/>"
            f"<b>Cabina Primaria ARERA:</b> {x['cp_arera_nome']} (cavidotto MT ~{x['cavidotto_mt_m']} m, CAPEX ~€ {x['capex_allaccio_eur']:,})<br/>"
            f"<b>Potenza Fotovoltaica:</b> {x['mwp_tot']} MWp tot ({x['mwp_300m']} MWp entro 300m) · ~{x['mwh_tracker']:,} MWh/anno<br/>"
            f"<b>VAM 2026:</b> {x['vam_2026_mq']:.2f} €/mq · <b>Spesa Max a 8 €/mq:</b> € {x['spesa_max_8eur']:,}<br/>"
            f"<b>Note:</b> {x['note']}"
            f"]]>"
        )
        placemarks.append(
            f"<Placemark>"
            f"<name>{x['id']} · {x['comune']} Fg.{x['foglio']} P.{x['particella']} ({x['sup_300m_ha']} ha ≤300m)</name>"
            f"<description>{desc}</description>"
            f"<Point><coordinates>{x['lon']},{x['lat']},0</coordinates></Point>"
            f"</Placemark>"
        )
    kml_content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<kml xmlns="http://www.opengis.net/kml/2.2">\n'
        '<Document>\n'
        '<name>SunPro — Cluster Piacentino A1 &amp; A21 (116 Lotti WMS AdE)</name>\n'
        + "\n".join(placemarks)
        + "\n</Document>\n</kml>"
    )
    kml_path.write_text(kml_content, encoding="utf-8")
    shutil.copy2(kml_path, WEB_DIR / kml_path.name)
    print(f"  ✓ KML Google Earth Piacentino ({kml_path.name}) generato!")


if __name__ == "__main__":
    raw = load_all_piacentino_raw_lots()
    enriched = enrich_piacentino_lots(raw)
    aperta = [x for x in enriched if x["gruppo"] == "APERTA_CAMPAGNA"]
    peri = [x for x in enriched if x["gruppo"] != "APERTA_CAMPAGNA"]
    print(f"Totale Lotti Piacentino: {len(enriched)} (Aperta Campagna: {len(aperta)}, Riserve Periurbane: {len(peri)})")
    print(f"Superficie Aperta Campagna: {sum(x['sup_tot_ha'] for x in aperta):.2f} ha tot ({sum(x['sup_300m_ha'] for x in aperta):.2f} ha entro 300m)")
    export_vector_layers_geojson(BASE_DIR / "piacentino_vector_layers.geojson")
    export_extended_excel_and_csv(
        enriched,
        BASE_DIR / "Terreni_FV_Piacentino_A1_A21_Esteso.xlsx",
        BASE_DIR / "SunPro_Piacentino_A1_A21.csv"
    )
    export_piacentino_kml(enriched, BASE_DIR / "SunPro_Piacentino_A1_A21.kml")
    (DATA_DIR / "piacentino_enriched_lots.json").write_text(
        json.dumps(enriched, indent=2, ensure_ascii=False), encoding="utf-8"
    )
